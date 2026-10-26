"""Session insights beyond the answers: camera analysis (emotions, presence), interview integrity, and CV profiles."""
import io
import logging
import re
from typing import Dict, List, Optional

import process_interview as pi

logger = logging.getLogger("evalbot.insights")

EMOTIONS = ["neutral", "happy", "surprised", "sad", "angry", "fearful", "disgusted"]


# --------------------------------------------------------------------------------------
# Camera analysis (all numbers were computed in the candidate's browser)
# --------------------------------------------------------------------------------------
def camera_summary(answers: List[Dict]) -> Optional[Dict]:
    pres = [(a["idx"], a["presence"]) for a in answers if (a.get("presence") or {}).get("samples")]
    if not pres:
        return None
    total = sum(p["samples"] for _, p in pres)

    def wavg(key):
        vals = [(p.get(key), p["samples"]) for _, p in pres if p.get(key) is not None]
        n = sum(w for _, w in vals)
        return round(sum(v * w for v, w in vals) / n, 3) if n else None

    emo_tot = {e: 0.0 for e in EMOTIONS}
    emo_n = 0
    timeline = []
    for idx, p in pres:
        em = p.get("emotions")
        if em:
            for e in EMOTIONS:
                emo_tot[e] += float(em.get(e, 0)) * p["samples"]
            emo_n += p["samples"]
        for seg in p.get("timeline") or []:
            timeline.append({"answer": idx + 1, "t": seg.get("t", 0), "dominant": seg.get("dominant"),
                             "probs": {e: round(float((seg.get("probs") or {}).get(e, 0)), 3) for e in EMOTIONS}})
    emotions = {e: round(v / emo_n, 3) for e, v in emo_tot.items()} if emo_n else None
    return {
        "samples": total, "answers": len(pres),
        "face_visible": wavg("face_visible"), "eye_contact": wavg("eye_contact"), "smile": wavg("smile"),
        "head_movement": wavg("head_movement"),
        "emotions": emotions, "dominant_emotion": max(emotions, key=emotions.get) if emotions else None,
        "timeline": timeline[:400],
    }


# --------------------------------------------------------------------------------------
# Integrity: events become review flags with a simple, explainable score
# --------------------------------------------------------------------------------------
# type -> (penalty per event, cap)
PENALTIES = {
    "multiple_faces": (15, 30), "second_voice": (15, 30), "voice_without_lips": (12, 24), "tab_hidden": (8, 32),
    "reading_pattern": (8, 16), "looking_away": (3, 15), "no_face": (4, 16), "fullscreen_exit": (3, 9),
}


def integrity(answers: List[Dict], camera: Optional[Dict], camera_required: bool) -> Dict:
    events: List[Dict] = []
    for a in answers:
        p = a.get("presence") or {}
        for e in p.get("events") or []:
            if e.get("type") in PENALTIES:
                events.append({"type": e["type"], "answer": a["idx"] + 1, "t": round(float(e.get("t", 0)), 1),
                               "dur": round(float(e["dur"]), 1) if e.get("dur") is not None else None})
        if (a.get("voice") or {}).get("second_voice"):
            events.append({"type": "second_voice", "answer": a["idx"] + 1, "t": 0, "dur": None})
    counts: Dict[str, int] = {}
    for e in events:
        counts[e["type"]] = counts.get(e["type"], 0) + 1
    penalty = sum(min(PENALTIES[k][0] * n, PENALTIES[k][1]) for k, n in counts.items())
    notes = []
    if camera:
        away = 1 - (camera.get("eye_contact") or 0)
        if away > 0.35:
            penalty += 15
            notes.append("away_high")
        elif away > 0.2:
            penalty += 6
            notes.append("away_some")
        fv = camera.get("face_visible") or 0
        if fv < 0.6:
            penalty += 15
            notes.append("face_often_missing")
        elif fv < 0.85:
            penalty += 6
            notes.append("face_sometimes_missing")
    elif camera_required:
        penalty += 25
        notes.append("camera_off")
    score = max(0, 100 - penalty)
    return {"score": score, "level": "low" if score >= 80 else "medium" if score >= 60 else "high",
            "counts": counts, "events": events[:200], "notes": notes, "camera": bool(camera)}


# --------------------------------------------------------------------------------------
# CV profiles
# --------------------------------------------------------------------------------------
CV_SCHEMA = {
    "type": "object",
    "properties": {
        "name": {"type": "string"},
        "headline": {"type": "string", "description": "one line, e.g. 'Data engineer with 2 years of Spark and Kafka'"},
        "years_experience": {"type": "number"},
        "skills": {"type": "array", "items": {"type": "string"}, "description": "technical skills and tools, short names"},
        "soft_skills": {"type": "array", "items": {"type": "string"}},
        "experience": {"type": "array", "items": {"type": "object", "properties": {
            "title": {"type": "string"}, "company": {"type": "string"}, "period": {"type": "string"},
            "summary": {"type": "string"}}, "required": ["title", "company", "period", "summary"]}},
        "education": {"type": "array", "items": {"type": "object", "properties": {
            "degree": {"type": "string"}, "institution": {"type": "string"}, "year": {"type": "string"}},
            "required": ["degree", "institution", "year"]}},
        "projects": {"type": "array", "items": {"type": "object", "properties": {
            "name": {"type": "string"}, "summary": {"type": "string"}, "tech": {"type": "array", "items": {"type": "string"}}},
            "required": ["name", "summary", "tech"]}},
        "certifications": {"type": "array", "items": {"type": "string"}},
        "languages": {"type": "array", "items": {"type": "string"}, "description": "spoken languages only"},
    },
    "required": ["name", "headline", "years_experience", "skills", "soft_skills", "experience", "education", "projects",
                 "certifications", "languages"],
}


def cv_text_from_file(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
    elif name.endswith(".docx"):
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="ignore")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()[:20000]


def parse_cv(text: str) -> Optional[Dict]:
    prompt = f"""Extract a structured profile from this CV / resume. Use only what the CV says; do not invent anything.
Keep the original language of names and titles. Skills are short tool or technology names (e.g. "Python", "Apache Spark").

CV:
{text[:15000]}
"""
    return pi._gemini_json(prompt, CV_SCHEMA, temperature=0.1)


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9+#]", "", s.lower().replace("apache ", ""))


def skill_match(job_skills: List[str], cv: Optional[Dict], cv_text: str = "") -> Optional[Dict]:
    if not job_skills or not cv:
        return None
    have = [_norm(x) for x in (cv.get("skills") or []) + [t for p in cv.get("projects") or [] for t in p.get("tech") or []]]
    blob = _norm(cv_text or "")
    matched, missing = [], []
    for s in job_skills:
        k = _norm(s)
        ok = bool(k) and (any(k == h or (len(k) > 2 and len(h) > 2 and (k in h or h in k)) for h in have if h) or (len(k) > 2 and k in blob))
        (matched if ok else missing).append(s)
    return {"matched": matched, "missing": missing, "percent": round(100 * len(matched) / len(job_skills))}
